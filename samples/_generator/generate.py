#!/usr/bin/env python3
"""Regenerate the binary sample documents (PDF / XLSX / CSV / PNG) in samples/.

Text samples (01, 03, 07, 08) are hand-written and live directly in samples/.

Usage (deps isolated in a venv, never in the project root):
    python3 -m venv /path/to/venv
    /path/to/venv/bin/pip install reportlab openpyxl pillow
    /path/to/venv/bin/python samples/_generator/generate.py

All companies, people, addresses, part numbers and contract numbers are fictional.
Phone numbers use the 555-01xx fictional range. "Today" for every sample is 2026-10-02.
"""
from __future__ import annotations

import csv
import math
import os
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

SAMPLES = Path(__file__).resolve().parent.parent
random.seed(20261002)

# --------------------------------------------------------------------------- styles
BASE = ParagraphStyle("base", fontName="Helvetica", fontSize=8, leading=10, alignment=TA_LEFT)
SMALL = ParagraphStyle("small", parent=BASE, fontSize=6.5, leading=8)
BOLD = ParagraphStyle("bold", parent=BASE, fontName="Helvetica-Bold")
MONO = ParagraphStyle("mono", parent=BASE, fontName="Courier", fontSize=8, leading=9.5)
TITLE = ParagraphStyle("title", parent=BASE, fontName="Helvetica-Bold", fontSize=16, leading=19)
CO = ParagraphStyle("co", parent=BASE, fontName="Helvetica-Bold", fontSize=12, leading=14)
TC = ParagraphStyle("tc", parent=BASE, fontSize=6.6, leading=8.2)
TCH = ParagraphStyle("tch", parent=BASE, fontName="Helvetica-Bold", fontSize=7.5, leading=9.5)

GRID = colors.HexColor("#555555")
SHADE = colors.HexColor("#E4E8EE")


def P(text: str, style=BASE) -> Paragraph:
    return Paragraph(text, style)


def boxed(rows, col_widths, header_rows=0, style_extra=None):
    t = Table(rows, colWidths=col_widths)
    cmds = [
        ("BOX", (0, 0), (-1, -1), 0.8, GRID),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, GRID),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    if header_rows:
        cmds.append(("BACKGROUND", (0, 0), (-1, header_rows - 1), SHADE))
    if style_extra:
        cmds.extend(style_extra)
    t.setStyle(TableStyle(cmds))
    return t


def build_with_page_count(path: Path, story, left: str, form: str, **doc_kwargs):
    """Two-pass build so the footer can say 'Page x of N'."""

    class _Doc(SimpleDocTemplate):
        pass

    # pass 1: count pages
    tmp = _Doc(str(path), pagesize=letter, **doc_kwargs)
    pages = {"n": 0}

    def count(canv, doc):
        pages["n"] = doc.page

    tmp.build(list(story), onFirstPage=count, onLaterPages=count)
    total = pages["n"]

    def _footer(canv, doc):
        canv.saveState()
        canv.setFont("Helvetica", 6.5)
        canv.setFillColor(colors.HexColor("#333333"))
        canv.drawString(0.5 * inch, 0.38 * inch, left)
        canv.drawCentredString(4.25 * inch, 0.38 * inch, form)
        canv.drawRightString(8.0 * inch, 0.38 * inch, f"Page {doc.page} of {total}")
        canv.restoreState()

    doc = _Doc(str(path), pagesize=letter, **doc_kwargs)
    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


def logo_block(name: str, tagline: str, addr: str, mark: str, color: str = "#1F3A5F"):
    mark_tbl = Table(
        [[P(f'<font color="white"><b>{mark}</b></font>', ParagraphStyle("m", parent=CO, fontSize=13, alignment=1))]],
        colWidths=[0.55 * inch],
        rowHeights=[0.5 * inch],
    )
    mark_tbl.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(color)),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    txt = [P(name, CO), P(tagline, SMALL), P(addr, BASE)]
    t = Table([[mark_tbl, txt]], colWidths=[0.65 * inch, 3.2 * inch])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    return t


# --------------------------------------------------------------------------- 02 aerospace PO
ARCTURUS_TERMS = [
    ("ACCEPTANCE", "Supplier's acknowledgment, commencement of work, or shipment of any goods constitutes acceptance of this Purchase Order and these Terms. Additional or different terms proposed by Supplier are rejected unless expressly accepted in writing by an authorized Arcturus Buyer."),
    ("PRICE AND INVOICING", "Prices shall not exceed those stated on the face of this order. Where price is shown as TBD, Supplier shall confirm price in writing prior to shipment. Invoices must reference PO number and line number; one invoice per shipment."),
    ("DELIVERY", "Time is of the essence. Supplier shall notify Buyer immediately of any anticipated delay. Early shipments greater than five (5) days require Buyer approval. Arcturus may return over-shipments at Supplier's expense."),
    ("QUALITY REQUIREMENTS", "Supplier shall maintain a quality system compliant with ISO 9001 or AS9100 as applicable. Quality clauses referenced on the face of this order are incorporated by reference per Arcturus QAP-100. Raw material shall be traceable to the producing mill heat/lot."),
    ("CERTIFICATIONS", "Each shipment shall include a Certificate of Conformance signed by an authorized representative and legible copies of mill test reports showing chemical and mechanical properties to the specification and revision cited."),
    ("SPECIALTY METALS / DFARS", "When this order supports a U.S. Government prime contract, Supplier shall comply with DFARS 252.225-7009 and 252.225-7008 as applicable, and shall provide country-of-melt on certifications. Clauses are flowed down in full."),
    ("COUNTERFEIT PREVENTION", "Supplier shall not deliver counterfeit or suspect counterfeit material and shall flow down this requirement to sub-tier suppliers per DFARS 252.246-7007."),
    ("RIGHT OF ACCESS", "Arcturus, its customers, and regulatory authorities shall have right of access to Supplier's facilities and records applicable to this order."),
    ("FOREIGN OBJECT DEBRIS", "Supplier shall maintain an FOD prevention program. Material shall be packaged to prevent damage, corrosion and contamination in transit."),
    ("WARRANTY", "Supplier warrants that all goods conform to specifications, are free from defects in material and workmanship, and are fit for the purpose intended for a period of twelve (12) months from acceptance."),
    ("CHANGES", "Buyer may make changes within the general scope of this order by written change order. Supplier shall assert any claim for adjustment within twenty (20) days of receipt of the change."),
    ("TERMINATION", "Buyer may terminate this order in whole or in part for convenience or for Supplier default in accordance with Arcturus Standard Terms, Form AA-PUR-012."),
    ("GOVERNING LAW", "This order shall be governed by the laws of the State of Kansas, excluding its conflict of laws provisions and the U.N. Convention on Contracts for the International Sale of Goods."),
]


def make_02():
    path = SAMPLES / "02-aerospace-po.pdf"
    story = []

    hdr_right = boxed(
        [
            [P("<b>PURCHASE ORDER</b>", TITLE), ""],
            [P("PO Number", BOLD), P("<b>4500118274</b>", ParagraphStyle("pn", parent=BOLD, fontSize=10))],
            [P("Revision", BOLD), P("0")],
            [P("PO Date", BOLD), P("10/01/2026")],
            [P("Supplier No.", BOLD), P("V-20417")],
            [P("Buyer", BOLD), P("Dana Whitfield")],
            [P("Page", BOLD), P("1 of 2")],
        ],
        [1.1 * inch, 2.0 * inch],
        style_extra=[("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (1, 0), SHADE)],
    )
    head = Table(
        [
            [
                logo_block(
                    "ARCTURUS AEROSTRUCTURES, INC.",
                    "WING &amp; EMPENNAGE ASSEMBLIES  |  AS9100D / NADCAP CHEM PROCESSING",
                    "4410 N. Webb Road<br/>Wichita, KS 67226<br/>Tel (316) 555-0100",
                    "AA",
                ),
                hdr_right,
            ]
        ],
        colWidths=[4.4 * inch, 3.1 * inch],
    )
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 8)]

    addr = boxed(
        [
            [P("<b>SUPPLIER / VENDOR</b>"), P("<b>SHIP TO</b>"), P("<b>BILL TO</b>")],
            [
                P("Nox Metals<br/>Detroit, MI<br/>Attn: Inside Sales<br/>sales@noxmetals.co"),
                P("Arcturus Aerostructures, Inc.<br/>Plant 3 - Receiving, Dock 7<br/>4500 N. Webb Road, Bldg 3<br/>Wichita, KS 67226<br/>Recv hrs 06:00-14:30 M-F"),
                P("Arcturus Aerostructures, Inc.<br/>Attn: Accounts Payable<br/>P.O. Box 7781<br/>Wichita, KS 67277<br/>ap-invoices@arcturus-aero.com"),
            ],
        ],
        [2.5 * inch] * 3,
        header_rows=1,
    )
    story += [addr, Spacer(1, 6)]

    terms = boxed(
        [
            [P("<b>Ship Via</b>"), P("<b>FOB</b>"), P("<b>Freight Terms</b>"), P("<b>Payment Terms</b>"), P("<b>Buyer Contact</b>")],
            [
                P("FEDEX FREIGHT PRIORITY"),
                P("ORIGIN"),
                P("PREPAY &amp; ADD"),
                P("NET 45"),
                P("Dana Whitfield, Sr. Buyer<br/>(316) 555-0129<br/>dwhitfield@arcturus-aero.com"),
            ],
            [P("<b>Prime Contract</b>"), P("<b>DPAS Rating</b>"), P("<b>Program</b>"), P("<b>Quality Clauses</b>"), P("<b>Confirm To</b>")],
            [P("FA8620-26-C-0041"), P("DO-A2"), P("AT-6 WING SPAR KITS"), P("Q01, Q05, Q12, Q31"), P("Buyer - email acknowledgment within 24 hrs")],
        ],
        [1.45 * inch, 0.95 * inch, 1.5 * inch, 1.4 * inch, 2.2 * inch],
        style_extra=[("BACKGROUND", (0, 0), (-1, 0), SHADE), ("BACKGROUND", (0, 2), (-1, 2), SHADE)],
    )
    story += [terms, Spacer(1, 8)]

    lines = [
        ("10", "RM-7075-T7351-0750",
         "PLATE, ALUMINUM 7075-T7351<br/>.750 THK X 24.00 X 36.00 IN<br/>SPEC: AMS-QQ-A-250/12<br/>SAW CUT TOL +.125/-0",
         "6", "EA"),
        ("20", "RM-7050-T7451-2000",
         "PLATE, AL 7050-T7451<br/>2.000 X 12.000 X 18.000 IN<br/>SPEC: AMS 4050<br/>DFARS 252.225-7009 COMPLIANT MATERIAL REQ'D",
         "4", "EA"),
        ("30", "RM-2024-T351-1250",
         "PLATE, AL 2024-T351<br/>1.250 X 10.00 X 20.00 IN<br/>SPEC: AMS-QQ-A-250/4",
         "8", "EA"),
        ("40", "RM-6061-T651-0500",
         "PLATE, AL 6061-T651  .500 X 48.0 X 96.0<br/>SPEC: AMS 4027 / ASTM B209<br/>FULL SHEET - DO NOT CUT",
         "2", "EA"),
    ]
    rows = [[P("<b>Ln</b>", TCH), P("<b>Item / Part No.</b>", TCH), P("<b>Description</b>", TCH),
             P("<b>Qty</b>", TCH), P("<b>UM</b>", TCH), P("<b>Date Req'd</b>", TCH),
             P("<b>Unit Price</b>", TCH), P("<b>Ext. Amount</b>", TCH)]]
    for ln, pn, desc, qty, um in lines:
        rows.append([P(ln), P(pn, MONO), P(desc), P(qty), P(um), P("10/05/2026"), P("TBD"), P("TBD")])
    rows.append(["", "", P("<b>*** PRICE TO BE CONFIRMED BY SUPPLIER PRIOR TO ACCEPTANCE ***</b>", SMALL), "", "", "", P("<b>Subtotal</b>"), P("TBD")])
    rows.append(["", "", "", "", "", "", P("<b>Freight</b>"), P("TBD")])
    rows.append(["", "", "", "", "", "", P("<b>PO TOTAL</b>"), P("<b>TBD</b>")])
    tbl = boxed(
        rows,
        [0.35 * inch, 1.45 * inch, 2.75 * inch, 0.4 * inch, 0.35 * inch, 0.75 * inch, 0.7 * inch, 0.75 * inch],
        header_rows=1,
        style_extra=[("ALIGN", (3, 1), (3, -1), "RIGHT"), ("SPAN", (2, 5), (5, 5))],
    )
    story += [tbl, Spacer(1, 8)]

    notes = [
        "1. ALL MATERIAL SHALL BE DFARS COMPLIANT PER DFARS 252.225-7009. COUNTRY OF MELT TO BE SHOWN ON CERTS.",
        "2. CERTIFIED MILL TEST REPORTS (MTRs) AND SIGNED C OF C REQUIRED WITH EACH SHIPMENT. REFERENCE PO AND LINE NUMBER ON ALL DOCUMENTS.",
        "3. MATERIAL MUST BE MERCURY FREE. NO WELD REPAIR.",
        "4. EACH PIECE TO BE INK MARKED WITH ALLOY, TEMPER, SPEC AND HEAT/LOT NO.",
        "5. NO PARTIAL SHIPMENTS WITHOUT BUYER APPROVAL. NO SUBSTITUTIONS.",
        "6. <b>EXPEDITE - MATERIAL REQUIRED ON DOCK 10/05/2026 TO SUPPORT LINE LOAD.</b> SUPPLIER TO CONFIRM PRICE AND DELIVERY BY RETURN EMAIL WITHIN 24 HRS.",
    ]
    story.append(
        boxed(
            [[P("<b>SPECIAL INSTRUCTIONS / NOTES</b>")], [P("<br/>".join(notes))]],
            [7.5 * inch],
            header_rows=1,
        )
    )
    story.append(Spacer(1, 10))
    sig = boxed(
        [
            [P("<b>Authorized By</b>"), P("<b>Title</b>"), P("<b>Date</b>"), P("<b>Supplier Acknowledgment (sign &amp; return)</b>")],
            [P("<i>D. Whitfield</i> (electronic)"), P("Senior Buyer, Raw Material"), P("10/01/2026"), P("__________________________  Date ________")],
        ],
        [1.6 * inch, 1.7 * inch, 0.9 * inch, 3.3 * inch],
        header_rows=1,
    )
    story += [KeepTogether([sig]), Spacer(1, 6)]
    story.append(P("THIS PURCHASE ORDER IS SUBJECT TO ARCTURUS STANDARD TERMS AND CONDITIONS OF PURCHASE (FORM AA-PUR-012 REV 2024-03), INCLUDED HEREIN AND AVAILABLE ON REQUEST. "
                   "ITEMS ON THIS ORDER MAY SUPPORT A DEFENSE PRIORITIES AND ALLOCATIONS SYSTEM (DPAS) RATED ORDER CERTIFIED FOR NATIONAL DEFENSE USE.", SMALL))

    story.append(PageBreak())
    story.append(P("ARCTURUS AEROSTRUCTURES, INC. - STANDARD TERMS AND CONDITIONS OF PURCHASE", ParagraphStyle("t2", parent=BOLD, fontSize=10, leading=13)))
    story.append(P("Form AA-PUR-012 Rev 2024-03  |  Applies to PO 4500118274", SMALL))
    story.append(Spacer(1, 6))
    left, right = [], []
    for i, (h, body) in enumerate(ARCTURUS_TERMS, 1):
        (left if i <= 7 else right).append(P(f"<b>{i}. {h}.</b> {body}", TC))
        (left if i <= 7 else right).append(Spacer(1, 4))
    cols = Table([[left, right]], colWidths=[3.7 * inch, 3.7 * inch])
    cols.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(cols)

    build_with_page_count(
        path, story,
        left="Printed 10/01/2026 15:22 CDT by DWHITFIELD  |  SAP ME21N",
        form="PO 4500118274 Rev 0",
        leftMargin=0.5 * inch, rightMargin=0.5 * inch, topMargin=0.5 * inch, bottomMargin=0.6 * inch,
        title="Purchase Order 4500118274", author="Arcturus Aerostructures, Inc.",
    )
    return path


# --------------------------------------------------------------------------- 05 RFQ w/ out-of-catalog line
def make_05():
    path = SAMPLES / "05-rfq-out-of-catalog.pdf"
    story = []
    hdr_right = boxed(
        [
            [P("<b>REQUEST FOR QUOTATION</b>", ParagraphStyle("rt", parent=TITLE, fontSize=13, leading=16)), ""],
            [P("RFQ No.", BOLD), P("<b>RFQ-26-0912</b>")],
            [P("Issue Date", BOLD), P("09/30/2026")],
            [P("Quote Due", BOLD), P("10/07/2026  5:00 PM CT")],
            [P("Buyer", BOLD), P("Kevin Marsh")],
        ],
        [1.0 * inch, 2.1 * inch],
        style_extra=[("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (1, 0), SHADE)],
    )
    head = Table(
        [[logo_block("RIDGELINE DEFENSE COMPONENTS", "PRECISION MACHINED ASSEMBLIES  |  ISO 9001:2015  |  ITAR REGISTERED",
                     "1850 Sparkman Dr NW<br/>Huntsville, AL 35816<br/>(256) 555-0150", "RD"), hdr_right]],
        colWidths=[4.4 * inch, 3.1 * inch],
    )
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 8)]

    story.append(boxed(
        [
            [P("<b>TO (SUPPLIER)</b>"), P("<b>DELIVER TO</b>"), P("<b>BUYER CONTACT</b>")],
            [
                P("Nox Metals<br/>Detroit, MI<br/>sales@noxmetals.co"),
                P("Ridgeline Defense Components<br/>Shipping &amp; Receiving<br/>1850 Sparkman Dr NW<br/>Huntsville, AL 35816"),
                P("Kevin Marsh, Commodity Buyer<br/>Direct: (256) 555-0173<br/>kmarsh@ridgeline-dc.com"),
            ],
        ],
        [2.5 * inch] * 3, header_rows=1,
    ))
    story.append(Spacer(1, 6))
    story.append(boxed(
        [
            [P("<b>Delivery Req'd</b>"), P("<b>FOB</b>"), P("<b>Ship Via</b>"), P("<b>Terms</b>"), P("<b>Quote Validity</b>")],
            [P("2 WEEKS ARO"), P("DESTINATION"), P("BEST WAY - PREPAID"), P("NET 30"), P("60 DAYS MIN.")],
        ],
        [1.5 * inch] * 5, header_rows=1,
    ))
    story.append(Spacer(1, 8))

    lines = [
        ("1", "RDC-10442", "PLATE, AL 6061-T651<br/>.625 X 12.00 X 12.00<br/>AMS 4027", "250", "EA"),
        ("2", "RDC-20871", "BAR, ROUND, CRES 17-4PH, COND H1025<br/>1.500 DIA X 12.00 LG<br/>AMS 5643", "20", "EA"),
        ("3", "RDC-30115", "PLATE, AL 7075-T7351<br/>1.000 X 6.00 X 12.00<br/>AMS 4078", "40", "EA"),
        ("4", "RDC-30220", "BAR, RD, AL 7075-T651<br/>3.000 DIA X 18.00 LG<br/>AMS 4122 / AMS-QQ-A-225/9", "12", "EA"),
    ]
    rows = [[P(h, TCH) for h in ["<b>Item</b>", "<b>Part No.</b>", "<b>Description / Specification</b>", "<b>Qty</b>", "<b>UM</b>",
                                   "<b>Unit Price</b>", "<b>Lead Time</b>", "<b>No Quote</b>"]]]
    for it, pn, desc, qty, um in lines:
        rows.append([P(it), P(pn, MONO), P(desc), P(qty), P(um), P(""), P(""), P("[  ]")])
    rows.append(["", "", P("<b>RFQ - PRICING REQUESTED. SUPPLIER TO COMPLETE SHADED COLUMNS.</b>", SMALL), "", "", "", "", ""])
    story.append(boxed(
        rows,
        [0.4 * inch, 0.95 * inch, 2.85 * inch, 0.45 * inch, 0.35 * inch, 0.85 * inch, 0.85 * inch, 0.8 * inch],
        header_rows=1,
        style_extra=[("BACKGROUND", (5, 1), (6, 4), colors.HexColor("#FFF6D5")), ("SPAN", (2, 5), (7, 5)),
                     ("ALIGN", (3, 1), (3, -1), "RIGHT")],
    ))
    story.append(Spacer(1, 8))
    notes = [
        "A. ALL MATERIAL MUST BE DFARS COMPLIANT (MELTED IN USA OR QUALIFYING COUNTRY). STATE COUNTRY OF MELT IN QUOTE.",
        "B. MILL TEST REPORTS REQUIRED. CERTS MUST SHOW SPEC AND REVISION.",
        "C. QUOTE EACH LINE SEPARATELY. PARTIAL QUOTES ACCEPTED - MARK 'NO QUOTE' FOR ANY LINE NOT QUOTED.",
        "D. INDICATE IF MATERIAL IS FROM STOCK OR MILL. STATE MINIMUM ORDER / CUT CHARGES IF ANY.",
        "E. THIS RFQ DOES NOT CONTAIN TECHNICAL DATA. DO NOT RETURN DRAWINGS.",
    ]
    story.append(boxed([[P("<b>QUOTE INSTRUCTIONS</b>")], [P("<br/>".join(notes))]], [7.5 * inch], header_rows=1))
    story.append(Spacer(1, 8))
    story.append(P(
        "Issuance of this RFQ does not obligate Ridgeline Defense Components to award a purchase order or pay any cost incurred in preparation of a quotation. "
        "Any resulting order is subject to Ridgeline Terms &amp; Conditions RDC-QP-7.4 and flowdown clauses applicable to U.S. Government subcontracts, including "
        "FAR 52.204-21, DFARS 252.204-7012, DFARS 252.225-7009 and DFARS 252.246-7007. Supplier must be registered in SAM.gov. "
        "Return completed quotation to the buyer by the due date shown above.", SMALL))
    build_with_page_count(
        path, story,
        left="RDC Purchasing  |  Generated by Epicor Kinetic 2025.2  |  09/30/2026 11:08",
        form="RFQ-26-0912",
        leftMargin=0.5 * inch, rightMargin=0.5 * inch, topMargin=0.5 * inch, bottomMargin=0.6 * inch,
        title="RFQ-26-0912", author="Ridgeline Defense Components",
    )
    return path


# --------------------------------------------------------------------------- 04 BOM xlsx + csv
BOM_META = [
    ["BOM EXPORT - Assy HR-AMR-200-A00  Rev C  (AMR Chassis, Gen 2)"],
    ["Exported 10/02/2026 08:15 by pnatarajan  |  Source: Arena PLM  |  Build Qty: 2  |  Need Date: 10/16/2026"],
    ["RFQ to: Nox Metals <sales@noxmetals.co>  |  Contact: Priya Natarajan, Sr. Buyer, Halvorsen Robotics, 616-555-0109, pnatarajan@halvorsenrobotics.com"],
    ["Ship to: Halvorsen Robotics, 3350 Kraft Ave SE, Grand Rapids, MI 49512  |  Certs required on all raw material"],
    [],
]
BOM_HEADER = ["Find #", "Part Number", "Descr.", "Matl / Spec", "Raw Stock Size", "Units", "Qty/Assy", "Ext Qty", "Make/Buy", "Notes"]
BOM_ROWS = [
    [10, "HR-1001-A", "BASE PLATE", "AL 6061-T651", "1.00 x 12.00 x 18.00", "IN", 2, 4, "MAKE", ""],
    [20, "HR-1002-B", "GANTRY SIDE PLATE", "AL6061T651", ".750 X 8 X 30", "IN", 4, 8, "MAKE", "L/R hand, same blank"],
    [30, "HR-1003", "TOOLING PLATE", "MIC6", "1/2 x 18 x 24", "IN", 1, 2, "MAKE", "cast plate, must be flat"],
    [40, "HR-1010", "END CAP", "7075-T6", "25.4 x 150 x 300", "mm", 3, 6, "MAKE", "metric dwg - see HR-1010 sht 2"],
    [50, "HR-1011", "LINEAR SHAFT BLANK", "6061-T6511 RD", "1.25 DIA x 36", "IN", 2, 4, "MAKE", ""],
    [60, "HR-2001", "SHCS M6x20 CL12.9", "STL BLK OX", "-", "EA", 24, 48, "BUY", "stock room"],
    [70, "HR-1020", "CABLE TRAY COVER", "5052-H32", ".090 x 10 x 40", "IN", 1, 2, "MAKE", "brake formed"],
]


def make_04():
    xlsx = SAMPLES / "04-bom-export.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM_HR-AMR-200"
    for r in BOM_META:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=12)
    for c in ("A2", "A3", "A4"):
        ws[c].font = Font(italic=True, size=9, color="444444")
    ws.append(BOM_HEADER)
    hdr_row = ws.max_row
    thin = Side(style="thin", color="999999")
    for cell in ws[hdr_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for r in BOM_ROWS:
        ws.append(r)
        for cell in ws[ws.max_row]:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if r[8] == "BUY":
            for cell in ws[ws.max_row]:
                cell.fill = PatternFill("solid", fgColor="EDEDED")
    widths = [8, 13, 22, 16, 20, 7, 9, 8, 10, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    wb.save(xlsx)

    csv_path = SAMPLES / "04-bom-export.csv"
    with open(csv_path, "w", newline="") as f:
        w = csv.writer(f)
        for r in BOM_META:
            w.writerow(r)
        w.writerow(BOM_HEADER)
        for r in BOM_ROWS:
            w.writerow(r)
    return xlsx, csv_path


# --------------------------------------------------------------------------- 06 handwritten note (PNG + image-only PDF)
HAND_FONTS = [
    "/System/Library/Fonts/Supplemental/Bradley Hand Bold.ttf",
    "/System/Library/Fonts/Noteworthy.ttc",
    "/System/Library/Fonts/MarkerFelt.ttc",
]


def _font(paths, size):
    for p in paths:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size)
            except OSError:
                continue
    return ImageFont.load_default(size)


def _hand_line(draw, xy, text, font, fill, jitter=2):
    """Draw text glyph by glyph with a little baseline wobble so it reads as handwriting."""
    x, y = xy
    for ch in text:
        dy = random.uniform(-jitter, jitter)
        draw.text((x, y + dy), ch, font=font, fill=fill)
        x += draw.textlength(ch, font=font) + random.uniform(-0.6, 0.9)
    return x


def make_06():
    W, H = 1600, 2050
    paper = Image.new("RGB", (W, H), (250, 247, 236))
    px = paper.load()
    # paper grain
    for _ in range(140000):
        x, y = random.randrange(W), random.randrange(H)
        v = random.randint(-14, 6)
        r, g, b = px[x, y]
        px[x, y] = (max(0, r + v), max(0, g + v), max(0, b + v))
    d = ImageDraw.Draw(paper)

    # printed memo-pad header
    pf = _font(["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/System/Library/Fonts/Helvetica.ttc"], 46)
    pf2 = _font(["/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Helvetica.ttc"], 26)
    d.rectangle([0, 0, W, 190], fill=(214, 226, 240))
    d.text((90, 40), "VUKOVICH TOOL & DIE, INC.", font=pf, fill=(30, 55, 95))
    d.text((92, 108), "23900 Ryan Rd  •  Warren, MI 48091  •  (586) 555-0140", font=pf2, fill=(30, 55, 95))
    d.text((92, 142), "Dies • Fixtures • Gauges • Wire EDM", font=pf2, fill=(60, 80, 115))
    # ruled lines + margin
    for y in range(300, H - 60, 96):
        d.line([(60, y), (W - 60, y)], fill=(160, 190, 220), width=2)
    d.line([(170, 200), (170, H)], fill=(225, 150, 150), width=3)

    ink = (28, 42, 120)
    big = _font(HAND_FONTS, 66)
    mid = _font(HAND_FONTS, 58)
    base_y = 300 - 66  # write sitting on the rules

    def row(i):
        return base_y + 96 * i

    _hand_line(d, (1220, row(0)), "10/2", mid, ink)
    _hand_line(d, (200, row(1)), "NOX - pls quote:", big, ink)
    _hand_line(d, (230, row(2)), "(3)  6061-T651   2 x 4 x 6", mid, ink)
    _hand_line(d, (230, row(3)), "(12) 7075-T651  .375 x 3 x 5", mid, ink)
    _hand_line(d, (230, row(4)), "(1)  5083   1/4 x 48 x 96", mid, ink)
    end = _hand_line(d, (230, row(5)), "(2)  2024   1 x 6 x 6", mid, ink)
    # scribble it out
    y0 = row(5) + 40
    pts = []
    x = 215
    while x < end + 20:
        pts.append((x, y0 + random.randint(-14, 14)))
        x += 22
    d.line(pts, fill=ink, width=6)
    d.line([(215, y0 - 4), (end + 25, y0 + 8)], fill=ink, width=5)
    _hand_line(d, (end + 50, row(5)), "no", mid, ink)

    _hand_line(d, (200, row(7)), "need Mon AM !!", big, ink)
    d.line([(200, row(7) + 92), (720, row(7) + 86)], fill=ink, width=5)
    _hand_line(d, (200, row(8)), "certs w/ matl", mid, ink)
    _hand_line(d, (200, row(10)), "Tony V.   586-555-0147", mid, ink)
    _hand_line(d, (200, row(11)), "PO to follow", mid, ink)

    # coffee ring
    ring = Image.new("L", (W, H), 0)
    rd = ImageDraw.Draw(ring)
    cx, cy, rr = 1290, 1620, 190
    rd.ellipse([cx - rr, cy - rr, cx + rr, cy + rr], outline=90, width=16)
    ring = ring.filter(ImageFilter.GaussianBlur(5))
    stain = Image.new("RGB", (W, H), (176, 128, 70))
    paper = Image.composite(stain, paper, ring.point(lambda v: int(v * 0.55)))

    # photographed: rotate on a desk, uneven lighting, slight blur
    desk = Image.new("RGB", (W + 260, H + 260), (92, 84, 74))
    rot = paper.rotate(-2.6, resample=Image.BICUBIC, expand=True, fillcolor=(92, 84, 74))
    desk.paste(rot, (130 - (rot.width - W) // 2, 120 - (rot.height - H) // 2))
    gw, gh = desk.size
    grad = Image.linear_gradient("L").resize((gw * 2, gh * 2)).rotate(35, resample=Image.BICUBIC)
    grad = grad.crop((gw // 2, gh // 2, gw // 2 + gw, gh // 2 + gh)).filter(ImageFilter.GaussianBlur(40))
    shade = Image.new("RGB", desk.size, (0, 0, 0))
    desk = Image.composite(shade, desk, grad.point(lambda v: int(v * 0.22)))
    desk = desk.filter(ImageFilter.GaussianBlur(1.1))
    noise = Image.effect_noise(desk.size, 10).convert("RGB")
    desk = Image.blend(desk, noise, 0.06)
    out = desk.resize((desk.width * 3 // 4, desk.height * 3 // 4), Image.LANCZOS)
    png = SAMPLES / "06-scanned-note.png"
    out.save(png, optimize=True)

    # image-only PDF: same note run through an office copier/scanner (grayscale, skewed, speckled)
    scan = paper.convert("L").rotate(1.4, resample=Image.BICUBIC, expand=True, fillcolor=255)
    scan = scan.point(lambda v: 255 if v > 205 else max(0, v - 25))
    sp = scan.load()
    for _ in range(9000):
        x, y = random.randrange(scan.width), random.randrange(scan.height)
        sp[x, y] = random.randint(0, 90)
    scan = scan.filter(ImageFilter.GaussianBlur(0.7)).resize((scan.width // 2, scan.height // 2), Image.LANCZOS)
    tmp_jpg = SAMPLES / "_generator" / "_scan_tmp.jpg"
    scan.save(tmp_jpg, quality=55)
    from reportlab.pdfgen import canvas as rl_canvas

    pdf = SAMPLES / "06-scanned-note.pdf"
    c = rl_canvas.Canvas(str(pdf), pagesize=letter)
    c.setTitle("SCAN0042")
    c.setAuthor("RICOH MP C3004")
    pw, ph = letter
    ratio = min((pw - 30) / scan.width, (ph - 30) / scan.height)
    iw, ih = scan.width * ratio, scan.height * ratio
    c.drawImage(str(tmp_jpg), (pw - iw) / 2, (ph - ih) / 2, iw, ih)
    c.showPage()
    c.save()
    tmp_jpg.unlink()
    return png, pdf


# --------------------------------------------------------------------------- 09 launch-vehicle PO
# Fictional company. Clause set follows conventions common to space-industry raw-material POs:
# AS9100 flow-down, AS9102 FAI, export-control notice, DFARS specialty metals (nickel alloy line only),
# counterfeit prevention, FOD, heat-lot traceability.
CORVANE_TERMS = [
    ("ACCEPTANCE", "Supplier's written acknowledgment, commencement of performance, or shipment constitutes acceptance of this Purchase Order, including all referenced Supplier Quality Clauses. Supplier's quotation terms do not apply."),
    ("QUALITY SYSTEM", "Supplier shall maintain a quality management system compliant with AS9100 or ISO 9001 and shall flow down the requirements of this order, including key characteristics and special requirements, to sub-tier sources."),
    ("RAW MATERIAL CERTIFICATION", "Each shipment shall include a Certificate of Conformance and the producing mill's certified test report showing chemistry and mechanical properties to the specification and revision cited, with heat/lot number traceable to each piece."),
    ("FIRST ARTICLE", "Where SQ-014 is invoked, the first shipment of each part number requires First Article Inspection documentation per AS9102. Cut-to-size raw stock: dimensional FAI only."),
    ("SPECIALTY METALS", "Lines identified as specialty metals (including nickel-base alloys and titanium) shall comply with DFARS 252.225-7009. Supplier shall state country of melt on the certification."),
    ("COUNTERFEIT PREVENTION", "Supplier shall procure only from the original mill or its authorized distributors and shall maintain a counterfeit-material prevention program consistent with AS6174."),
    ("EXPORT CONTROL", "Drawings and specifications referenced on this order may contain technical data controlled under the ITAR (22 CFR 120-130) or EAR. Supplier shall not export, re-export or disclose such data to foreign persons without authorization."),
    ("FOREIGN OBJECT DEBRIS", "Material shall be free of FOD, packaged to prevent damage and contamination, and plate shall be interleaved or film-protected on both faces."),
    ("NONCONFORMANCE", "Supplier shall notify Buyer within 24 hours of discovering any nonconformance on material already shipped. Supplier may not disposition nonconforming material as use-as-is."),
    ("RIGHT OF ACCESS", "Corvane, its customers and regulatory agencies shall have right of access to Supplier's facilities and records pertaining to this order."),
    ("RECORD RETENTION", "Quality records for this order shall be retained for not less than ten (10) years after final payment."),
    ("GOVERNING LAW", "This order is governed by the laws of the State of Colorado, excluding conflict of laws provisions."),
]


def corvane_logo():
    """Original vector mark: an ascent trajectory cutting through an orbit ring, plus wordmark."""
    from reportlab.graphics.shapes import Drawing, Group, Line, PolyLine, Polygon, String, Wedge

    ink = colors.HexColor("#14213D")
    accent = colors.HexColor("#C8553D")
    d = Drawing(3.3 * inch, 0.62 * inch)

    ring = Group()
    cx, cy, r = 22, 22, 17
    for start, extent in ((200, 250),):
        w = Wedge(cx, cy, r, start, start + extent, radius1=r - 2.4, strokeColor=None, fillColor=ink)
        ring.add(w)
    d.add(ring)
    trajectory = PolyLine([4, 6, 14, 12, 24, 22, 33, 36, 38, 44], strokeColor=accent, strokeWidth=2.6, strokeLineCap=1)
    d.add(trajectory)
    d.add(Polygon([38, 44, 41.5, 42.5, 40.5, 47.5], fillColor=accent, strokeColor=accent, strokeWidth=1.2))

    d.add(String(50, 23, "CORVANE", fontName="Helvetica-Bold", fontSize=19, fillColor=ink))
    d.add(Line(50, 17, 199, 17, strokeColor=accent, strokeWidth=1))
    d.add(String(50, 6, "LAUNCH  SYSTEMS", fontName="Helvetica", fontSize=7.5, fillColor=ink))
    return d


def make_09():
    path = SAMPLES / "09-launch-vehicle-po.pdf"
    story = []

    hdr_right = boxed(
        [
            [P("<b>PURCHASE ORDER</b>", TITLE), ""],
            [P("PO Number", BOLD), P("<b>CLS-PO-26-03318</b>", ParagraphStyle("pn9", parent=BOLD, fontSize=10))],
            [P("Revision", BOLD), P("A")],
            [P("Order Date", BOLD), P("09/29/2026")],
            [P("Rev Date", BOLD), P("10/01/2026")],
            [P("Supplier ID", BOLD), P("S-08812")],
            [P("Buyer", BOLD), P("Priya Raman")],
        ],
        [1.1 * inch, 2.0 * inch],
        style_extra=[("SPAN", (0, 0), (1, 0)), ("BACKGROUND", (0, 0), (1, 0), SHADE)],
    )
    head = Table(
        [[
            [
                corvane_logo(),
                Spacer(1, 3),
                P("LAUNCH VEHICLES  |  AS9100D CERTIFIED  |  CAGE 9XK41", SMALL),
                P("11850 E. Corvane Way, Englewood, CO 80112  |  Tel (720) 555-0160", BASE),
            ],
            hdr_right,
        ]],
        colWidths=[4.4 * inch, 3.1 * inch],
    )
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [head, Spacer(1, 8)]

    addr = boxed(
        [
            [P("<b>SUPPLIER</b>"), P("<b>SHIP TO</b>"), P("<b>BILL TO</b>")],
            [
                P("Nox Metals<br/>Detroit, MI<br/>Attn: Inside Sales<br/>sales@noxmetals.co"),
                P("Corvane Launch Systems, Inc.<br/>Building 4 Receiving, Door 12<br/>11920 E. Corvane Way<br/>Englewood, CO 80112<br/>Recv 07:00-15:00 M-F"),
                P("Corvane Launch Systems, Inc.<br/>Accounts Payable<br/>11850 E. Corvane Way<br/>Englewood, CO 80112<br/>ap@corvanelaunch.com"),
            ],
        ],
        [2.5 * inch] * 3,
        header_rows=1,
    )
    story += [addr, Spacer(1, 6)]

    terms = boxed(
        [
            [P("<b>Ship Via</b>"), P("<b>FOB</b>"), P("<b>Payment Terms</b>"), P("<b>Program / Charge No.</b>"), P("<b>Buyer Contact</b>")],
            [P("UPS FREIGHT - CORVANE ACCT 4471-0093"), P("ORIGIN"), P("NET 30"), P("CL-2 STAGE 2 / 26-2210-MFG"),
             P("Priya Raman, Commodity Buyer<br/>(720) 555-0164<br/>praman@corvanelaunch.com")],
            [P("<b>Supplier Quality Clauses</b>"), P("<b>Export Class.</b>"), P("<b>Need By</b>"), P("<b>Govt Contract</b>"), P("<b>Acknowledge To</b>")],
            [P("SQ-001, SQ-014, SQ-022, SQ-031"), P("EAR99 (material)"), P("10/16/2026"), P("FA8811-25-C-0007 (line 40 only)"), P("Buyer, within 2 business days")],
        ],
        [1.55 * inch, 1.05 * inch, 0.9 * inch, 1.75 * inch, 2.25 * inch],
        style_extra=[("BACKGROUND", (0, 0), (-1, 0), SHADE), ("BACKGROUND", (0, 2), (-1, 2), SHADE)],
    )
    story += [terms, Spacer(1, 8)]

    lines = [
        ("10", "MS-2219-T87-1000",
         "PLATE, ALUMINUM 2219-T87<br/>1.000 X 48.00 X 96.00 IN<br/>SPEC: ASTM B209<br/>ULTRASONIC INSPECTED, CLASS A",
         "3", "EA"),
        ("20", "MS-7075-T7351-1500",
         "PLATE, AL 7075-T7351<br/>1.500 X 12.00 X 24.00 IN<br/>SPEC: AMS 4078<br/>SAW CUT TOL +.060/-0<br/><i>REV A: QTY WAS 10</i>",
         "12", "EA"),
        ("30", "MS-6061-T6511-RB3000",
         "BAR, RD, AL 6061-T6511<br/>3.000 DIA X 12.00 LG<br/>SPEC: ASTM B221",
         "20", "EA"),
        ("40", "MS-718-RB1250",
         "BAR, RD, NICKEL ALLOY 718<br/>1.250 DIA X 36.00 LG, SOLN TREATED<br/>SPEC: AMS 5662<br/>DFARS 252.225-7009 SPECIALTY METAL",
         "6", "EA"),
        ("50", "MS-7050-PL3000",
         "PLATE, AL 7050, TEMPER PER DWG 400-1182<br/>3.000 X 18.00 X 30.00 IN",
         "2", "EA"),
    ]
    rows = [[P("<b>Ln</b>", TCH), P("<b>Corvane P/N</b>", TCH), P("<b>Description</b>", TCH),
             P("<b>Qty</b>", TCH), P("<b>UM</b>", TCH), P("<b>Need By</b>", TCH),
             P("<b>Unit Price</b>", TCH), P("<b>Ext.</b>", TCH)]]
    for ln, pn, desc, qty, um in lines:
        rows.append([P(ln), P(pn, MONO), P(desc), P(qty), P(um), P("10/16/2026"), P("QUOTE"), P("QUOTE")])
    rows.append(["", "", P("<b>*** PRICING PER SUPPLIER QUOTE - CONFIRM ON ACKNOWLEDGMENT ***</b>", SMALL), "", "", "", P("<b>Total</b>"), P("TBD")])
    tbl = boxed(
        rows,
        [0.35 * inch, 1.5 * inch, 2.8 * inch, 0.4 * inch, 0.35 * inch, 0.75 * inch, 0.65 * inch, 0.7 * inch],
        header_rows=1,
        style_extra=[("ALIGN", (3, 1), (3, -1), "RIGHT"), ("SPAN", (2, 6), (5, 6))],
    )
    story += [tbl, Spacer(1, 8)]

    notes = [
        "1. REVISION A (10/01/2026): LINE 20 QTY CHANGED FROM 10 TO 12. ALL OTHER TERMS UNCHANGED.",
        "2. CERTIFIED MILL TEST REPORTS AND C OF C REQUIRED WITH SHIPMENT. HEAT/LOT NO. MARKED ON EACH PIECE (SQ-022).",
        "3. SQ-014: FIRST ARTICLE (AS9102, DIMENSIONAL ONLY) REQUIRED ON FIRST SHIPMENT OF EACH CORVANE P/N.",
        "4. LINE 40 SUPPORTS U.S. GOVERNMENT CONTRACT FA8811-25-C-0007. DFARS 252.225-7009 APPLIES TO LINE 40 ONLY. STATE COUNTRY OF MELT.",
        "5. LINE 50 TEMPER PER CORVANE DWG 400-1182 (REV C). DRAWING MAY CONTAIN EXPORT-CONTROLLED TECHNICAL DATA; REQUEST VIA BUYER.",
        "6. PLATE TO BE FILM-PROTECTED BOTH FACES. NO MIXED HEATS WITHIN A LINE. PARTIAL SHIPMENTS ACCEPTABLE.",
    ]
    story.append(boxed([[P("<b>NOTES</b>")], [P("<br/>".join(notes))]], [7.5 * inch], header_rows=1))
    story.append(Spacer(1, 10))
    sig = boxed(
        [
            [P("<b>Issued By</b>"), P("<b>Title</b>"), P("<b>Date</b>"), P("<b>Supplier Acknowledgment</b>")],
            [P("<i>P. Raman</i> (electronic)"), P("Commodity Buyer, Metals"), P("10/01/2026"), P("__________________________  Date ________")],
        ],
        [1.6 * inch, 1.7 * inch, 0.9 * inch, 3.3 * inch],
        header_rows=1,
    )
    story += [KeepTogether([sig]), Spacer(1, 6)]
    story.append(P("THIS ORDER IS SUBJECT TO CORVANE TERMS AND CONDITIONS OF PURCHASE (CLS-QP-110 REV F) AND THE SUPPLIER QUALITY CLAUSES LISTED ABOVE, WHICH ARE INCORPORATED BY REFERENCE.", SMALL))

    story.append(PageBreak())
    story.append(P("CORVANE LAUNCH SYSTEMS, INC. - TERMS AND CONDITIONS OF PURCHASE", ParagraphStyle("t9", parent=BOLD, fontSize=10, leading=13)))
    story.append(P("CLS-QP-110 Rev F  |  Applies to CLS-PO-26-03318 Rev A", SMALL))
    story.append(Spacer(1, 6))
    left, right = [], []
    for i, (h, body) in enumerate(CORVANE_TERMS, 1):
        (left if i <= 6 else right).append(P(f"<b>{i}. {h}.</b> {body}", TC))
        (left if i <= 6 else right).append(Spacer(1, 4))
    cols = Table([[left, right]], colWidths=[3.7 * inch, 3.7 * inch])
    cols.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(cols)

    build_with_page_count(
        path, story,
        left="Generated 10/01/2026 09:47 MDT  |  Corvane ERP / PUR-2",
        form="CLS-PO-26-03318 Rev A",
        leftMargin=0.5 * inch, rightMargin=0.5 * inch, topMargin=0.5 * inch, bottomMargin=0.6 * inch,
        title="Purchase Order CLS-PO-26-03318 Rev A", author="Corvane Launch Systems, Inc.",
    )
    return path


if __name__ == "__main__":
    for f in (make_02(), make_05(), *make_04(), *make_06(), make_09()):
        print("wrote", Path(f).relative_to(SAMPLES.parent))
